from openpyxl import Workbook


def generar_reporte(
    validos,
    rechazados,
    colisiones
):

    libro = Workbook()

    # Hoja válidos
    hoja_validos = libro.active
    hoja_validos.title = "Carga_Sistema"

    hoja_validos.append([
        "Consecutivo",
        "Tipo",
        "Perfil",
        "Correo",
        "Telefono",
        "Requiere_Aprobacion"
    ])

    for registro in validos:

        hoja_validos.append([
            registro.get("consecutivo"),
            registro.get("tipo"),
            registro.get("perfil_sistema"),
            registro.get("correo_corporativo"),
            registro.get("telefono"),
            registro.get("requiere_aprobacion")
        ])

    # Hoja rechazados
    hoja_rechazados = libro.create_sheet(
        "Rechazados"
    )

    hoja_rechazados.append([
        "Errores"
    ])

    for registro in rechazados:

        hoja_rechazados.append([
            "; ".join(
                registro["errores"]
            )
        ])

    # Hoja colisiones
    
    hoja_colisiones = libro.create_sheet(
        "Colisiones_Correo"
    )

    hoja_colisiones.append([
        "Consecutivo",
        "Base",
        "Correo_Generado"
    ])

    for colision in colisiones:

        hoja_colisiones.append([
            colision["consecutivo"],
            colision["base"],
            colision["correo_generado"]
        ])

    # Hoja resumen
    hoja_resumen = libro.create_sheet(
        "Resumen"
    )

    hoja_resumen.append([
        "Concepto",
        "Cantidad"
    ])

    hoja_resumen.append([
        "Validos",
        len(validos)
    ])

    hoja_resumen.append([
        "Rechazados",
        len(rechazados)
    ])

    hoja_resumen.append([
        "Total",
        len(validos) + len(rechazados)
    ])

    libro.save(
        "salida/reporte_novedades.xlsx"
    )