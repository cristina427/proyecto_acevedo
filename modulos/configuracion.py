from openpyxl import load_workbook
from modulos.excepciones import ErrorConfiguracion


def cargar_configuracion():
    try:
        libro = load_workbook("config/parametros.xlsx")
        hojas_requeridas = [
            "Parametros",
            "Cargos",
            "Centros_Costo",
            "Matriz_Roles",
            "Campos_Obligatorios"
        ]

        for hoja in hojas_requeridas:
            if hoja not in libro.sheetnames:
                raise ErrorConfiguracion(
                    f"Falta la hoja: {hoja}"
                )
        hoja_parametros = libro["Parametros"]

        parametros = {}

        for clave, valor in hoja_parametros.iter_rows(min_row=2, values_only=True):
         parametros[clave] = valor

         hoja_campos = libro["Campos_Obligatorios"]

         campos_obligatorios = {}

         for tipo, campos in hoja_campos.iter_rows(
         min_row=2,
         values_only=True
):
          campos_obligatorios[tipo] = campos.split(";")
         print(campos_obligatorios) 
        return parametros, campos_obligatorios
        return parametros
    
    except FileNotFoundError:
     raise ErrorConfiguracion("No se encontró el archivo parametros.xlsx")