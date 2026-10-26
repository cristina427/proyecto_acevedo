from openpyxl import load_workbook
from modulos.excepciones import ErrorConfiguracion


def cargar_configuracion():
    try:
        libro = load_workbook("config/parametros.xlsx")
        hojas_requeridas = [
            "Parametros",
            "Cargos",
            "Centros_Costo",
            "Matriz_Roles",
            "Campos_Obligatorios"
        ]

        for hoja in hojas_requeridas:
            if hoja not in libro.sheetnames:
                raise ErrorConfiguracion(
                    f"Falta la hoja: {hoja}"
                )
        hoja_parametros = libro["Parametros"]

        parametros = {}

        for clave, valor in hoja_parametros.iter_rows(min_row=2, values_only=True):
         parametros[clave] = valor

        hoja_campos = libro["Campos_Obligatorios"]

        campos_obligatorios = {}

        for tipo, campos in hoja_campos.iter_rows(min_row=2, values_only=True):
         campos_obligatorios[tipo] = campos.split(";")

        hoja_centros = libro["Centros_Costo"]

        centros_costo = {}

        for codigo, nombre, area in hoja_centros.iter_rows(min_row=2, values_only=True):
         centros_costo[codigo] = {"nombre": nombre, "area": area}

         hoja_cargos = libro["Cargos"]

         cargos = {}

        for codigo_cargo, nombre_cargo, nivel, area in hoja_cargos.iter_rows(min_row=2, values_only=True):
         cargos[codigo_cargo] = {
        "nombre": nombre_cargo,
        "nivel": nivel,
        "area": area}

         hoja_roles = libro["Matriz_Roles"]

         matriz_roles = {}

        for codigo_cargo, area, perfil, aprobacion in hoja_roles.iter_rows(min_row=2, values_only=True):

         clave = f"{codigo_cargo}|{area}"

         matriz_roles[clave] = {"perfil": perfil, "requiere_aprobacion": aprobacion}

        return (parametros, campos_obligatorios, centros_costo, cargos, matriz_roles)
    
    except FileNotFoundError:
     raise ErrorConfiguracion("No se encontró el archivo parametros.xlsx")